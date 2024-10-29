import os
from fpdf import FPDF
from docx import Document

class PDFWriter:
    def convert_to_pdf(self, directory):
        if not os.path.isdir(directory):
            raise ValueError("La dirección proporcionada no es una carpeta válida.")

        for filename in os.listdir(directory):
            file_path = os.path.join(directory, filename)
            if os.path.isfile(file_path):
                if filename.endswith('.docx'):
                    pdf_filename = f"{os.path.splitext(filename)[0]}.pdf"
                    pdf_path = os.path.join(directory, pdf_filename)

                    self._convert_docx_to_pdf(file_path, pdf_path)

    def _convert_docx_to_pdf(self, docx_file, pdf_file):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=11)

        document = Document(docx_file)

        for paragraph in document.paragraphs:
            pdf.cell(200, 10, txt=paragraph.text.encode('latin-1', 'replace').decode('latin-1'), ln=True)

        pdf.output(pdf_file)
        print(f"Convertido: {docx_file} a {pdf_file}")